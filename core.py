# -*- coding: utf-8 -*-
"""
Ядро сверки деклараций НВОС с платежами за квартал.

Колонки находятся по заголовкам (не по позициям), лист с данными
выбирается автоматически:
  - платежи: лист, где есть «Сумма платёжного поручения» и «ИНН»;
  - декларации: лист, где есть «Номер декларации», «ИНН» и колонка (040)/(080).

Типы сверки:
  - sbros  — «Сумма платы за сбросы, всего (080)»
  - vybros — «Сумма платы за выбросы всего (040)»

Платёж засчитывается, если период = КВ.0{квартал}.* ;
периоды МС.* засчитываются с флагом «проверить»;
остальные (КВ другого квартала, ГД, СТРОЙКА...) не учитываются, но
показываются на листе платежей с пометкой.
"""

import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TOLERANCE = 1.0  # допуск в рублях

TYPES = {
    'sbros': {
        'label': 'сбросы',
        'title': 'Сбросы (080)',
        'decl_sum_marker': '(080)',
    },
    'vybros': {
        'label': 'выбросы',
        'title': 'Выбросы (040)',
        'decl_sum_marker': '(040)',
    },
}

GREEN = PatternFill('solid', fgColor='C6EFCE')
BLUE = PatternFill('solid', fgColor='BDD7EE')
YELLOW = PatternFill('solid', fgColor='FFEB9C')
RED = PatternFill('solid', fgColor='FFC7CE')
GRAY = PatternFill('solid', fgColor='EDEDED')
BOLD = Font(bold=True)
BOLD_BLUE = Font(bold=True, color='1F4E78')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

MONEY_FMT = '#,##0.00'


# ---------------------------------------------------------------- утилиты

def norm_header(v):
    if v is None:
        return ''
    return ' '.join(str(v).replace('ё', 'е').replace('Ё', 'Е').lower().split())


def inn_norm(v):
    """Возвращает (ключ_для_сопоставления, отображаемое_значение) или (None, None)."""
    if v is None:
        return None, None
    s = str(v).strip()
    if s.endswith('.0'):
        s = s[:-2]
    if not s.isdigit() or len(s) < 8:
        return None, None
    return s.lstrip('0') or None, s


def kpp_key(v):
    if v is None:
        return ''
    s = str(v).strip()
    if s.endswith('.0'):
        s = s[:-2]
    if not s.isdigit():
        return ''
    return s.lstrip('0')


def as_float(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def last_header_col(ws, header_row):
    last = 0
    for c in range(1, ws.max_column + 1):
        if ws.cell(header_row, c).value not in (None, ''):
            last = c
    return last


# ---------------------------------------------------------------- поиск таблиц

def _row_headers(ws, r):
    return {norm_header(ws.cell(r, c).value): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(r, c).value not in (None, '')}


def _find_col(headers, *, contains=None, equals=None, exclude=None):
    for h, c in headers.items():
        if exclude and any(e in h for e in exclude):
            continue
        if equals and h == equals:
            return c
        if contains and all(s in h for s in contains):
            return c
    return None


def find_payments_table(wb):
    """Ищет лист и шапку платежей. Возвращает dict или None."""
    for sn in wb.sheetnames:
        ws = wb[sn]
        for r in range(1, min(ws.max_row, 15) + 1):
            headers = _row_headers(ws, r)
            if not headers:
                continue
            col_sum = _find_col(headers, contains=['сумма', 'поручени'])
            col_inn = _find_col(headers, contains=['инн'])
            if not (col_sum and col_inn):
                continue
            return {
                'ws': ws,
                'sheet': sn,
                'header_row': r,
                'last_col': last_header_col(ws, r),
                'cols': {
                    'inn': col_inn,
                    'amount': col_sum,
                    'kpp': _find_col(headers, contains=['кпп']),
                    'period': _find_col(headers, contains=['период']),
                    'kbk': _find_col(headers, equals='кбк'),
                    'num': _find_col(headers, equals='номер'),
                    'date': _find_col(headers, equals='дата'),
                    'desc': _find_col(headers, equals='описание'),
                },
            }
    return None


def find_decl_table(wb):
    for sn in wb.sheetnames:
        ws = wb[sn]
        for r in range(1, min(ws.max_row, 15) + 1):
            headers = _row_headers(ws, r)
            if not headers:
                continue
            col_decl_num = _find_col(headers, contains=['номер декларации'])
            col_inn = _find_col(headers, equals='инн')
            col_040 = _find_col(headers, contains=['(040)'])
            col_080 = _find_col(headers, contains=['(080)'])
            # отбрасываем колонки врезанных «1/4» — нужна именно «сумма платы»
            if col_040 and 'сумма платы' not in [h for h, c in headers.items() if c == col_040][0]:
                col_040 = None
            if col_080 and 'сумма платы' not in [h for h, c in headers.items() if c == col_080][0]:
                col_080 = None
            if not (col_decl_num and col_inn and (col_040 or col_080)):
                continue
            return {
                'ws': ws,
                'sheet': sn,
                'header_row': r,
                'last_col': last_header_col(ws, r),
                'cols': {
                    'decl_num': col_decl_num,
                    'inn': col_inn,
                    'kpp': _find_col(headers, equals='кпп'),
                    'name': _find_col(headers, contains=['название организации']),
                    'subject': _find_col(headers, contains=['субъект']),
                    'status': _find_col(headers, equals='статус'),
                    'sum_vybros': col_040,
                    'sum_sbros': col_080,
                },
            }
    return None


# ---------------------------------------------------------------- загрузка

@dataclass
class DeclData:
    table: dict
    rows: list = field(default_factory=list)


def load_declarations(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    table = find_decl_table(wb)
    if table is None:
        raise ValueError(
            'В файле деклараций не найден лист с колонками «Номер декларации», '
            '«ИНН» и «Сумма платы ... (040)/(080)».')
    ws = table['ws']
    cols = table['cols']
    data = DeclData(table=table)
    for r in range(table['header_row'] + 1, ws.max_row + 1):
        key, disp = inn_norm(ws.cell(r, cols['inn']).value)
        name = ws.cell(r, cols['name']).value if cols['name'] else None
        if key is None and (name in (None, '', '-')):
            continue  # пустые строки и строка итогов
        decl_num_raw = ws.cell(r, cols['decl_num']).value
        try:
            decl_num = int(str(decl_num_raw).strip().rstrip('.0') or 0)
        except (ValueError, TypeError):
            decl_num = 0
        data.rows.append({
            'row': r,
            'inn_key': key,
            'inn': disp,
            'kpp': kpp_key(ws.cell(r, cols['kpp']).value) if cols['kpp'] else '',
            'name': name,
            'subject': ws.cell(r, cols['subject']).value if cols['subject'] else None,
            'status': ws.cell(r, cols['status']).value if cols['status'] else None,
            'decl_num': decl_num,
            'sum_vybros': as_float(ws.cell(r, cols['sum_vybros']).value) if cols['sum_vybros'] else 0.0,
            'sum_sbros': as_float(ws.cell(r, cols['sum_sbros']).value) if cols['sum_sbros'] else 0.0,
        })
    return data


@dataclass
class PayData:
    table: dict
    rows: list = field(default_factory=list)


def load_payments(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    table = find_payments_table(wb)
    if table is None:
        raise ValueError(
            'В файле платежей не найден лист с колонками «ИНН» и '
            '«Сумма платёжного поручения».')
    ws = table['ws']
    cols = table['cols']

    def cell(r, key):
        c = cols.get(key)
        return ws.cell(r, c).value if c else None

    data = PayData(table=table)
    for r in range(table['header_row'] + 1, ws.max_row + 1):
        amount = cell(r, 'amount')
        key, disp = inn_norm(cell(r, 'inn'))
        if not isinstance(amount, (int, float)):
            continue
        # period: None — колонки нет в файле; '' — колонка есть, значение пустое
        if cols.get('period') is None:
            period = None
        else:
            v = cell(r, 'period')
            period = str(v).strip() if v is not None else ''
        data.rows.append({
            'row': r,
            'inn_key': key,
            'inn': disp,
            'kpp': kpp_key(cell(r, 'kpp')),
            'amount': float(amount),
            'period': period,
            'num': cell(r, 'num'),
            'date': cell(r, 'date'),
            'desc': cell(r, 'desc'),
        })
    return data


# ---------------------------------------------------------------- сверка

def _period_year_ok(period, year):
    """Год периода: ищем 4-значный год в строке периода."""
    if year is None:
        return True
    return str(year) in re.findall(r'\d{4}', period)


def _range_year_ok(period, year):
    """Диапазон лет с двузначным хвостом: «ГД.2024-25», «КВ.4.24-25»."""
    if year is None:
        return False
    m = re.search(r'\d{2,4}\s*-\s*(\d{2})$', period)
    return bool(m) and 2000 + int(m.group(1)) == year


def classify_period(period, quarter, year=None):
    """quarter: 1..3 или 'year'; year: int или None (любой год).
    period: None — колонки периода нет (файл считается отфильтрованным).
    -> ('ok'|'flag'|'skip', note)"""
    if period is None:
        return 'ok', ''
    if period == '':
        return 'flag', 'период не указан — проверить'
    p = period.replace(',', '.')  # опечатки плательщиков: «ГД,00.2025»
    if quarter == 'year':
        # год: кварталы, годовой платёж, полугодия и доплата по итогам года
        if p.startswith(('КВ.', 'ГД.', 'ПГ.', 'ДОПЛ')):
            if _period_year_ok(p, year):
                return 'ok', ''
            if _range_year_ok(p, year):
                return 'flag', f'период {period} — диапазон лет, проверить'
            return 'skip', f'период {period} — другой год'
        if p.startswith('МС.') and _period_year_ok(p, year):
            return 'flag', f'период {period} — проверить'
        return 'skip', f'период {period} — не учтён'
    if p.startswith(f'КВ.0{quarter}.') or p.startswith(f'КВ.{quarter}.'):
        if _period_year_ok(p, year):
            return 'ok', ''
        return 'skip', f'период {period} — другой год'
    if p.startswith('МС.') and _period_year_ok(p, year):
        return 'flag', f'период {period} — проверить'
    return 'skip', f'период {period} — не учтён'


def reconcile_type(decl: DeclData, pays: PayData, sum_key: str, quarter, year=None):
    """Возвращает аннотации и статистику для одного типа платы.

    quarter: 1..3 — ожидание 1/4 годовой суммы; 'year' — вся годовая сумма.
    year: год периода платежа (int) или None — любой год.
    """
    divisor = 1 if quarter == 'year' else 4
    # --- платежи: учтённые по периоду ---
    counted = defaultdict(list)   # inn_key -> [pay_row_dict]
    pay_notes = {}                # row -> (counted: bool, note)
    skipped_periods = Counter()
    for p in pays.rows:
        verdict, note = classify_period(p['period'], quarter, year)
        if verdict == 'skip':
            pay_notes[p['row']] = (False, note)
            skipped_periods[p['period']] += 1
            continue
        pay_notes[p['row']] = (True, note)
        if p['inn_key']:
            counted[p['inn_key']].append(p)

    # --- декларации: дедупликация по ИНН (актуальна с max № декларации) ---
    by_inn = defaultdict(list)
    for d in decl.rows:
        if d['inn_key'] and d[sum_key] > 0:
            by_inn[d['inn_key']].append(d)
    chosen = {}  # inn_key -> row dict
    for inn, rows in by_inn.items():
        chosen[inn] = max(rows, key=lambda d: (d['decl_num'], d['row']))

    decl_kpps = defaultdict(set)
    for d in decl.rows:
        if d['inn_key'] and d['kpp']:
            decl_kpps[d['inn_key']].add(d['kpp'])

    # --- аннотация строк деклараций (всех, где сумма > 0) ---
    decl_ann = {}  # row -> dict
    stats = Counter()
    expected_total = 0.0
    paid_total = 0.0
    debtors = []  # для листа «Должники»: не оплачено + недоплата
    for d in decl.rows:
        if not (d[sum_key] > 0):
            continue
        plist = counted.get(d['inn_key'], []) if d['inn_key'] else []
        expected = round(d[sum_key] / divisor, 2)
        psum = round(sum(p['amount'] for p in plist), 2)
        delta = round(psum - expected, 2)

        flags = []
        is_dup = d['inn_key'] and len(by_inn[d['inn_key']]) > 1
        is_chosen = d['inn_key'] and chosen.get(d['inn_key']) is d
        if is_dup:
            flags.append(f'по ИНН {len(by_inn[d["inn_key"]])} декларации'
                         + ('' if is_chosen else ' — дубль, в сводке не учтена'))
        if any((p['period'] or '').startswith('МС.') for p in plist):
            flags.append('есть платежи с периодом МС — проверить')
        if any(p['period'] == '' for p in plist):
            flags.append('есть платежи без указания периода — проверить')
        pay_kpps = {p['kpp'] for p in plist if p['kpp']}
        if plist and d['kpp'] and d['kpp'] not in pay_kpps:
            flags.append('КПП платежа не совпадает с КПП декларации')

        if not plist:
            status, fill = '❌ Не оплачено', RED
        elif abs(delta) <= TOLERANCE:
            status, fill = '✅ Оплачено', GREEN
        elif delta > TOLERANCE:
            status, fill = '💰 Переплата', BLUE
        else:
            status, fill = '⚠️ Недоплата', YELLOW

        decl_ann[d['row']] = {
            'expected': expected, 'paid': psum, 'count': len(plist),
            'nums': ', '.join(str(p['num']) for p in plist if p['num'] is not None),
            'kpps': ', '.join(sorted(pay_kpps)),
            'periods': ', '.join(sorted({p['period'] for p in plist if p['period']})),
            'delta': delta if plist else None,
            'status': status, 'fill': fill, 'flags': ' | '.join(flags),
        }

        if is_chosen or not d['inn_key']:
            stats[status] += 1
            expected_total += expected
            paid_total += psum
            if status in ('❌ Не оплачено', '⚠️ Недоплата'):
                debtors.append({
                    'inn': d['inn'], 'name': d['name'], 'kpp': d['kpp'],
                    'subject': d['subject'], 'decl_status': d['status'],
                    'decl_num': d['decl_num'], 'year_sum': d[sum_key],
                    'expected': expected, 'paid': psum,
                    'rest': round(expected - psum, 2), 'status': status,
                })

    debtors.sort(key=lambda x: -x['rest'])

    # --- аннотация платежей ---
    inn_with_sum = set(by_inn)
    all_decl_inns = {d['inn_key'] for d in decl.rows if d['inn_key']}
    pay_ann = {}  # row -> dict
    no_decl_count = 0
    no_decl_sum = 0.0
    for p in pays.rows:
        is_counted, period_note = pay_notes[p['row']]
        notes = [period_note] if period_note else []
        d = chosen.get(p['inn_key']) if p['inn_key'] else None
        if p['inn_key'] is None:
            found, fill = 'нет', RED
            notes.append('ИНН не распознан')
        elif d is not None:
            found = 'да'
            kpp_ok = p['kpp'] and p['kpp'] in decl_kpps.get(p['inn_key'], set())
            if not kpp_ok:
                notes.append(f'КПП декл: {d["kpp"]}')
            fill = GREEN if kpp_ok else BLUE
        elif p['inn_key'] in all_decl_inns:
            found, fill = 'да', YELLOW
            notes.append('в декларации сумма по этому виду = 0')
            d = next(x for x in decl.rows if x['inn_key'] == p['inn_key'])
        else:
            found, fill = 'нет', RED
            notes.append('ИНН не найден в декларациях')
            if is_counted:
                no_decl_count += 1
                no_decl_sum += p['amount']
        if not is_counted:
            fill = GRAY
        pay_ann[p['row']] = {
            'found': found,
            'name': d['name'] if d else None,
            'kpp': d['kpp'] if d else None,
            'year_sum': d[sum_key] if d else None,
            'expected': round(d[sum_key] / divisor, 2) if d else None,
            'counted': 'да' if is_counted else 'нет',
            'notes': ' | '.join(n for n in notes if n),
            'fill': fill,
        }

    counted_sum = sum(p['amount'] for plist in counted.values() for p in plist)
    return {
        'decl_ann': decl_ann,
        'pay_ann': pay_ann,
        'debtors': debtors,
        'stats': {
            'decl_count': len(chosen),
            'paid_ok': stats['✅ Оплачено'],
            'overpaid': stats['💰 Переплата'],
            'underpaid': stats['⚠️ Недоплата'],
            'not_paid': stats['❌ Не оплачено'],
            'expected_total': round(expected_total, 2),
            'paid_total': round(paid_total, 2),
            'payments_counted_sum': round(counted_sum, 2),
            'payments_total': len(pays.rows),
            'payments_skipped': sum(skipped_periods.values()),
            'skipped_periods': dict(skipped_periods.most_common(10)),
            'no_decl_count': no_decl_count,
            'no_decl_sum': round(no_decl_sum, 2),
            'sheet': pays.table['sheet'],
        },
    }


# ---------------------------------------------------------------- сборка книги

def _copy_header(src_ws, header_row, dst_ws, last_col, extra, font_extra=BOLD_BLUE):
    for c in range(1, last_col + 1):
        cell = dst_ws.cell(1, c, src_ws.cell(header_row, c).value)
        cell.font = BOLD
    for i, h in enumerate(extra):
        cell = dst_ws.cell(1, last_col + 1 + i, h)
        cell.font = font_extra
        cell.alignment = CENTER


def _copy_row(src_ws, src_r, dst_ws, dst_r, last_col):
    for c in range(1, last_col + 1):
        dst_ws.cell(dst_r, c, src_ws.cell(src_r, c).value)


def _autosize(ws, max_width=55, sample_rows=400):
    widths = {}
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if r_idx > sample_rows:
            break
        for c_idx, v in enumerate(row, 1):
            if v is None:
                continue
            ln = max((len(line) for line in str(v).splitlines()), default=0)
            if ln > widths.get(c_idx, 0):
                widths[c_idx] = ln
    for c_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = min(w + 2, max_width)


def build_workbook(decl_path, payment_paths: dict, quarter, out_path, year=None):
    """
    payment_paths: {'sbros': path | None, 'vybros': path | None}
    quarter: 1..3 или 'year'; year: год периода платежа (int) или None
    Возвращает stats: {'sbros': {...}, 'vybros': {...}}
    """
    decl = load_declarations(decl_path)
    out = openpyxl.Workbook()
    out.remove(out.active)
    summary_ws = out.create_sheet('Сводка')
    all_stats = {}
    summary_blocks = []

    is_year = quarter == 'year'
    exp_label = 'Ожидание за год' if is_year else '1/4 расчётная'
    decl_extra = [exp_label, 'Σ платежей', 'Кол-во платежек', '№№ платёжек',
                  'КПП плательщика', 'Периоды платежей', 'Дельта (Σ − ожидание)',
                  'Статус сверки', 'Флаги']
    pay_extra = ['Декларация найдена', 'Название по декларации', 'КПП декл',
                 'Годовая сумма (декл)', exp_label, 'Учтён в сверке', 'Заметки']

    for type_key, cfg in TYPES.items():
        path = payment_paths.get(type_key)
        if not path:
            continue
        sum_col_key = 'sum_' + type_key
        if not decl.table['cols'].get(sum_col_key):
            marker = cfg['decl_sum_marker']
            raise ValueError(f'В файле деклараций нет колонки «Сумма платы ... {marker}», '
                             f'нужной для сверки типа «{cfg["label"]}».')
        pays = load_payments(path)
        res = reconcile_type(decl, pays, sum_col_key, quarter, year)
        all_stats[type_key] = res['stats']
        label = cfg['label']

        # --- Лист «Декларации (тип)» ---
        dws = out.create_sheet(f'Декларации — {label}')
        src = decl.table['ws']
        last_col = decl.table['last_col']
        _copy_header(src, decl.table['header_row'], dws, last_col, decl_extra)
        dst_r = 2
        money_cols = {last_col + 1, last_col + 2, last_col + 7}
        for d in decl.rows:
            ann = res['decl_ann'].get(d['row'])
            if ann is None:
                continue
            _copy_row(src, d['row'], dws, dst_r, last_col)
            vals = [ann['expected'], ann['paid'], ann['count'], ann['nums'],
                    ann['kpps'], ann['periods'], ann['delta'], ann['status'], ann['flags']]
            for i, v in enumerate(vals):
                cell = dws.cell(dst_r, last_col + 1 + i, v)
                if last_col + 1 + i in money_cols:
                    cell.number_format = MONEY_FMT
            for c in range(1, last_col + len(decl_extra) + 1):
                dws.cell(dst_r, c).fill = ann['fill']
            dst_r += 1
        dws.freeze_panes = 'A2'
        dws.auto_filter.ref = f'A1:{get_column_letter(last_col + len(decl_extra))}{dst_r - 1}'
        _autosize(dws)

        # --- Лист «Платежи (тип)» ---
        pws = out.create_sheet(f'Платежи — {label}')
        psrc = pays.table['ws']
        plast = pays.table['last_col']
        _copy_header(psrc, pays.table['header_row'], pws, plast, pay_extra)
        dst_r = 2
        pmoney = {plast + 4, plast + 5}
        for p in pays.rows:
            ann = res['pay_ann'][p['row']]
            _copy_row(psrc, p['row'], pws, dst_r, plast)
            vals = [ann['found'], ann['name'], ann['kpp'], ann['year_sum'],
                    ann['expected'], ann['counted'], ann['notes']]
            for i, v in enumerate(vals):
                cell = pws.cell(dst_r, plast + 1 + i, v)
                if plast + 1 + i in pmoney:
                    cell.number_format = MONEY_FMT
            for c in range(1, plast + len(pay_extra) + 1):
                pws.cell(dst_r, c).fill = ann['fill']
            dst_r += 1
        pws.freeze_panes = 'A2'
        pws.auto_filter.ref = f'A1:{get_column_letter(plast + len(pay_extra))}{dst_r - 1}'
        _autosize(pws)

        # --- Лист «Должники (тип)» ---
        if res['debtors']:
            dbws = out.create_sheet(f'Должники — {label}')
            headers = ['ИНН', 'Название', 'КПП', 'Субъект РФ', 'Статус декларации',
                       '№ декларации', 'Годовая сумма, ₽',
                       'Ожидание за год, ₽' if is_year else 'Ожидание 1/4, ₽',
                       'Оплачено, ₽', 'Остаток, ₽', 'Статус сверки']
            for i, h in enumerate(headers, 1):
                cell = dbws.cell(1, i, h)
                cell.font = BOLD_BLUE
                cell.alignment = CENTER
            for r_i, deb in enumerate(res['debtors'], 2):
                vals = [deb['inn'], deb['name'], deb['kpp'], deb['subject'],
                        deb['decl_status'], deb['decl_num'], deb['year_sum'],
                        deb['expected'], deb['paid'], deb['rest'], deb['status']]
                fill = RED if deb['status'] == '❌ Не оплачено' else YELLOW
                for c_i, v in enumerate(vals, 1):
                    cell = dbws.cell(r_i, c_i, v)
                    if c_i in (7, 8, 9, 10):
                        cell.number_format = MONEY_FMT
                    cell.fill = fill
            dbws.freeze_panes = 'A2'
            dbws.auto_filter.ref = f'A1:K{len(res["debtors"]) + 1}'
            _autosize(dbws)

        summary_blocks.append((cfg['title'], res['stats'], pays.table['sheet'],
                               Path(path).name))

    # --- Сводка ---
    ws = summary_ws
    period_title = 'ГОД' if is_year else f'{quarter} КВАРТАЛ'
    if year:
        period_title += f' {year}'
    else:
        period_title += ' (год не задан — любые)'
    ws.cell(1, 1, f'СВЕРКА ПЛАТЕЖЕЙ — {period_title}').font = Font(bold=True, size=14, color='1F4E78')
    ws.cell(2, 1, f'Файл деклараций: {Path(decl_path).name}'
                  f' (лист «{decl.table["sheet"]}», деклараций: {len(decl.rows)})')
    r = 4
    for title, st, sheet, fname in summary_blocks:
        ws.cell(r, 1, f'— {title} —').font = Font(bold=True, size=12)
        r += 1
        ws.cell(r, 1, f'Файл платежей: {fname} (лист «{sheet}»)')
        r += 1
        rows = [
            ('Деклараций с суммой > 0 (уник. ИНН)', st['decl_count'], None),
            ('✅ Оплачено', st['paid_ok'], GREEN),
            ('💰 Переплата', st['overpaid'], BLUE),
            ('⚠️ Недоплата', st['underpaid'], YELLOW),
            ('❌ Не оплачено', st['not_paid'], RED),
            ('Ожидаемая Σ за год, ₽' if is_year else 'Ожидаемая Σ за квартал (1/4), ₽',
             st['expected_total'], None),
            ('Поступило по декларациям, ₽', st['paid_total'], None),
            ('Учтено платежей всего, ₽', st['payments_counted_sum'], None),
            ('Платежей в файле, шт', st['payments_total'], None),
            ('Не учтено по периоду, шт', st['payments_skipped'], None),
            ('Платежи без декларации, шт', st['no_decl_count'], RED if st['no_decl_count'] else None),
            ('Платежи без декларации, ₽', st['no_decl_sum'], None),
        ]
        for label_, val, fill in rows:
            ca = ws.cell(r, 1, label_)
            cb = ws.cell(r, 2, val)
            if isinstance(val, float):
                cb.number_format = MONEY_FMT
            if fill:
                ca.fill = fill
                cb.fill = fill
            r += 1
        if st['skipped_periods']:
            ws.cell(r, 1, 'Неучтённые периоды: ' + ', '.join(
                f'{k or "(пусто)"}: {v}' for k, v in st['skipped_periods'].items()))
            r += 1
        r += 1
    ws.column_dimensions['A'].width = 42
    ws.column_dimensions['B'].width = 20

    out.save(out_path)
    return all_stats
